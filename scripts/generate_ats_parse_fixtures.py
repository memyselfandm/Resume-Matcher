"""Generate the synthetic ATS parse-check fixtures.

Run from ``apps/backend`` so the backend's dependencies are available::

    uv run python ../../scripts/generate_ats_parse_fixtures.py [fixture names...]

With fixture names, only those files are rewritten (Chromium is launched only
when a selected fixture needs it).
Set ``CHROMIUM_EXECUTABLE`` to reuse a locally installed Chromium binary.

Layout fixtures are rendered from small synthetic HTML documents with headless
Chromium (Playwright). Fixtures that need low-level PDF constructs (icon-font
glyphs, unmapped CID glyphs, a decompression bomb) are assembled byte by byte.
DOCX fixtures use python-docx. Every persona is fictional; contact details use
reserved example domains and 555 phone numbers.

The generated files are committed under
``apps/backend/tests/fixtures/ats_parse/`` so the default test suite never
needs Chromium. Re-run this script only when a fixture must change.
"""

from __future__ import annotations

import base64
import io
import json
import os
import sys
import zlib
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from playwright.sync_api import Page, sync_playwright

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO_ROOT / "apps" / "backend" / "tests" / "fixtures" / "ats_parse"

NAME = "Jordan Rivera"
EMAIL = "jordan.rivera@example.com"
PHONE = "(555) 010-4477"
LINKEDIN = "linkedin.com/in/jordan-rivera-example"
LOCATION = "Oakland, CA"

EXPERIENCE: list[dict[str, object]] = [
    {
        "title": "Senior Software Engineer",
        "company": "Northwind Analytics",
        "location": "San Francisco, CA",
        "years": "Jan 2021 - Present",
        "bullets": [
            "Led migration of 40 services to a shared deployment platform, reducing release time by 35%.",
            "Designed an event pipeline processing 2,000,000 records per day with 99.9% availability.",
            "Mentored 6 engineers.",
            "Automated compliance reporting, saving $120,000 per year in manual review effort.",
        ],
    },
    {
        "title": "Software Engineer",
        "company": "Bluebird Logistics",
        "location": "Oakland, CA",
        "years": "Jun 2017 - Dec 2020",
        "bullets": [
            "Built route optimization services used by 300 dispatchers across 12 regional hubs.",
            "Improved query latency by 60% by redesigning the reporting schema.",
            "Launched an internal API catalog.",
            "Implemented contract tests that reduced production incidents by 25%.",
        ],
    },
    {
        "title": "Junior Developer",
        "company": "Harbor Point Media",
        "location": "Vallejo, CA",
        "years": "Aug 2015 - May 2017",
        "bullets": [
            "Developed content publishing tools for a newsroom of 80 editors.",
            "Streamlined image processing jobs, cutting storage costs by 18%.",
            "Delivered 15 client dashboards on schedule.",
        ],
    },
]

EDUCATION = {
    "degree": "B.S. Computer Science",
    "institution": "Bay State University",
    "location": "Berkeley, CA",
    "years": "Aug 2011 - May 2015",
}

SWISS_EDUCATION_EXTRA = (
    ("Certificate, Data Engineering", "Harbor Extension School", "Oakland, CA", "Jan 2019 - Jun 2019"),
    ("Cloud Practitioner Credential", "Example Cloud Institute", "Remote", "Mar 2020 - Apr 2020"),
)


SUMMARY = (
    "Backend engineer with 9 years of experience building reliable data platforms, "
    "developer tooling, and logistics systems. Focused on measurable reliability and "
    "delivery improvements."
)

SKILLS = "Python, Go, PostgreSQL, Kafka, Kubernetes, Terraform, FastAPI, React"

BASE_CSS = """
@page { size: Letter; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; font-family: Helvetica, Arial, sans-serif; font-size: 10.5px;
       line-height: 1.35; color: #000; }
.page { padding: 38px 40px; }
h1 { font-size: 22px; margin: 0 0 4px 0; }
h3 { font-size: 12px; margin: 14px 0 6px 0; border-bottom: 1px solid #000; }
ul { margin: 3px 0 8px 0; padding-left: 16px; }
li { margin: 1px 0; }
p { margin: 0; }
"""


def _bullets(items: list[str]) -> str:
    return "<ul>" + "".join(f"<li>{item}</li>" for item in items) + "</ul>"


def clean_single_column_html() -> str:
    """Plain single-column resume with inline dates (no right alignment)."""
    jobs = "".join(
        f"<p><b>{job['title']}</b>, {job['company']}, {job['location']} "
        f"({job['years']})</p>{_bullets(job['bullets'])}"  # type: ignore[arg-type]
        for job in EXPERIENCE
    )
    return f"""<html><head><style>{BASE_CSS}</style></head><body><div class="page">
<h1>{NAME}</h1>
<p>{EMAIL} | {PHONE} | {LINKEDIN} | {LOCATION}</p>
<h3>Summary</h3><p>{SUMMARY}</p>
<h3>Experience</h3>{jobs}
<h3>Education</h3>
<p><b>{EDUCATION['degree']}</b>, {EDUCATION['institution']} ({EDUCATION['years']})</p>
<h3>Skills</h3><p>{SKILLS}</p>
</div></body></html>"""


def swiss_single_like_html() -> str:
    """Mirror swiss-single: uppercase headings, right-aligned dates and locations."""
    css = BASE_CSS + """
h3 { text-transform: uppercase; letter-spacing: 0.05em; }
.row { display: flex; justify-content: space-between; align-items: baseline; }
.row h4 { font-size: 11.5px; margin: 0; }
.sub { color: #333; }
.item { margin-bottom: 6px; }
.center { text-align: center; }
"""
    jobs = "".join(
        f"""<div class="item"><div class="row"><h4>{job['title']}</h4>
<span>{job['years']}</span></div>
<div class="row sub"><span>{job['company']}</span><span>{job['location']}</span></div>
{_bullets(job['bullets'])}</div>"""  # type: ignore[arg-type]
        for job in EXPERIENCE
    )
    education = "".join(
        f"""<div class="item"><div class="row"><h4>{degree}</h4><span>{years}</span></div>
<div class="row sub"><span>{institution}</span><span>{location}</span></div></div>"""
        for degree, institution, location, years in (
            (EDUCATION["degree"], EDUCATION["institution"], EDUCATION["location"], EDUCATION["years"]),
            *SWISS_EDUCATION_EXTRA,
        )
    )
    return f"""<html><head><style>{css}</style></head><body><div class="page">
<div class="center"><h1 style="text-transform: uppercase">{NAME}</h1>
<p>{EMAIL} &nbsp; {PHONE} &nbsp; {LINKEDIN} &nbsp; {LOCATION}</p></div>
<h3>Summary</h3><p style="text-align: justify">{SUMMARY}</p>
<h3>Experience</h3>{jobs}
<h3>Education</h3>{education}
<h3>Skills</h3><p>{SKILLS}</p>
</div></body></html>"""


def _two_column_html(main_fr: int, side_fr: int, gap_px: int) -> str:
    css = BASE_CSS + f"""
.grid {{ display: grid; grid-template-columns: minmax(0, {main_fr}fr) minmax(0, {side_fr}fr);
         gap: {gap_px}px; align-items: start; }}
h3 {{ text-transform: uppercase; }}
"""
    jobs = "".join(
        f"<p><b>{job['title']}</b></p><p>{job['company']} | {job['years']}</p>"
        f"{_bullets(job['bullets'])}"  # type: ignore[arg-type]
        for job in EXPERIENCE
    )
    side = f"""<h3>Contact</h3><p>{EMAIL}</p><p>{PHONE}</p><p>{LINKEDIN}</p><p>{LOCATION}</p>
<h3>Education</h3><p>{EDUCATION['degree']}</p><p>{EDUCATION['institution']}</p>
<p>{EDUCATION['years']}</p>
<h3>Skills</h3><p>Python</p><p>Go</p><p>PostgreSQL</p><p>Kafka</p><p>Kubernetes</p>
<p>Terraform</p><p>FastAPI</p><p>React</p>
<h3>Languages</h3><p>English (native)</p><p>Spanish (professional)</p>"""
    return f"""<html><head><style>{css}</style></head><body><div class="page">
<h1 style="text-align:center; text-transform: uppercase">{NAME}</h1>
<div class="grid"><div><h3>Summary</h3><p>{SUMMARY}</p><h3>Experience</h3>{jobs}</div>
<div>{side}</div></div>
</div></body></html>"""


def two_column_html() -> str:
    """swiss-two-column geometry: 65:35 grid with the default 16px gap."""
    return _two_column_html(65, 35, 16)


def balanced_two_column_html() -> str:
    """Two equal columns: multi-column, but neither side is a narrow sidebar."""
    return _two_column_html(50, 50, 24)


def spanish_html() -> str:
    """Spanish-language resume with diacritics in names and places."""
    return f"""<html lang="es"><head><style>{BASE_CSS}</style></head><body><div class="page">
<h1>José Núñez Peña</h1>
<p>jose.nunez@example.com | +34 555 010 223 | linkedin.com/in/jose-nunez-example | Málaga, España</p>
<h3>Resumen</h3><p>Ingeniero de software con ocho años de experiencia en el desarrollo de
plataformas de datos y servicios para la logística. Me enfoco en la fiabilidad y en la mejora
continua de los procesos del equipo.</p>
<h3>Experiencia</h3>
<p><b>Ingeniero de Software Sénior</b>, Compañía Andaluza de Datos (2019 - 2024)</p>
<ul><li>Dirigí la migración de 30 servicios a una plataforma común y reduje el tiempo de
publicación en un 40%.</li><li>Diseñé un sistema de informes para los equipos de operaciones
de la región y de las oficinas centrales.</li></ul>
<p><b>Ingeniero de Software</b>, Logística del Sur (2015 - 2019)</p>
<ul><li>Construí servicios de optimización de rutas para 200 conductores.</li>
<li>Mejoré la latencia de las consultas en un 50% con un nuevo esquema.</li></ul>
<h3>Educación</h3><p>Grado en Ingeniería Informática, Universidad de Ejemplo (2011 - 2015)</p>
<h3>Habilidades</h3><p>Python, Go, PostgreSQL, Kafka, Kubernetes, Terraform</p>
</div></body></html>"""


def long_document_html() -> str:
    """Twelve pages of resume-like content, for the page-cap check."""
    pages = "".join(
        f'<div class="page" style="page-break-after: always"><h3>Project Log {index}</h3>'
        f"<p>{SUMMARY}</p>{_bullets(EXPERIENCE[0]['bullets'])}</div>"  # type: ignore[arg-type]
        for index in range(1, 13)
    )
    return f"<html><head><style>{BASE_CSS}</style></head><body>{pages}</body></html>"


def swiss_single_source() -> dict[str, object]:
    """ResumeData-shaped payload rendered by ``swiss_single_like_html``."""
    education = [
        {
            "id": index,
            "degree": degree,
            "institution": institution,
            "years": years,
        }
        for index, (degree, institution, _location, years) in enumerate(
            (
                (EDUCATION["degree"], EDUCATION["institution"], EDUCATION["location"], EDUCATION["years"]),
                *SWISS_EDUCATION_EXTRA,
            )
        )
    ]
    return {
        "personalInfo": {
            "name": NAME,
            "email": EMAIL,
            "phone": PHONE,
            "linkedin": LINKEDIN,
            "location": LOCATION,
        },
        "summary": SUMMARY,
        "workExperience": [
            {
                "id": index,
                "title": job["title"],
                "company": job["company"],
                "location": job["location"],
                "years": job["years"],
                "description": job["bullets"],
            }
            for index, job in enumerate(EXPERIENCE)
        ],
        "education": education,
        "additional": {"technicalSkills": [skill.strip() for skill in SKILLS.split(",")]},
        "sectionMeta": [
            {"id": "summary", "key": "summary", "displayName": "Summary", "isVisible": True, "order": 1},
            {"id": "workExperience", "key": "workExperience", "displayName": "Experience", "isVisible": True, "order": 2},
            {"id": "education", "key": "education", "displayName": "Education", "isVisible": True, "order": 3},
            {"id": "additional", "key": "additional", "displayName": "Skills", "isVisible": True, "order": 4},
        ],
    }


def _render_pdf(page: Page, html: str) -> bytes:
    page.set_content(html, wait_until="load")
    return page.pdf(format="Letter", print_background=True, prefer_css_page_size=True)


def _image_only_pdf(page: Page) -> bytes:
    """Screenshot a rendered resume and print it back as a single raster image."""
    page.set_viewport_size({"width": 816, "height": 1056})
    page.set_content(clean_single_column_html(), wait_until="load")
    png = page.screenshot(full_page=False, type="png")
    encoded = base64.b64encode(png).decode("ascii")
    html = (
        "<html><head><style>@page { size: Letter; margin: 0; } body { margin: 0; }"
        "img { width: 8.5in; height: 11in; display: block; }</style></head>"
        f'<body><img src="data:image/png;base64,{encoded}"></body></html>'
    )
    return _render_pdf(page, html)


def _assemble_pdf(objects: list[bytes]) -> bytes:
    """Serialize numbered objects (1-based, object 1 is the catalog) into a PDF."""
    output = io.BytesIO()
    output.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_offset = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return output.getvalue()


def _stream(data: bytes, extra: bytes = b"") -> bytes:
    return (
        b"<< /Length " + str(len(data)).encode() + extra + b" >>\nstream\n"
        + data + b"\nendstream"
    )


def _text_ops(lines: list[tuple[str, int, int, str]]) -> bytes:
    """Build content-stream text operators from (font, x, y, literal) tuples."""
    parts = []
    for font, x, y, literal in lines:
        parts.append(f"BT /{font} 11 Tf {x} {y} Td {literal} Tj ET")
    return "\n".join(parts).encode("latin-1")


def _pdf_literal(text: str) -> str:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"({escaped})"


_BODY_LINES = [
    "Summary",
    "Backend engineer with 9 years of experience building data platforms.",
    "Experience",
    "Senior Software Engineer, Northwind Analytics, Jan 2021 - Present",
    "Led migration of 40 services, reducing release time by 35%.",
    "Designed an event pipeline processing 2,000,000 records per day.",
    "Education",
    "B.S. Computer Science, Bay State University, May 2015",
    "Skills",
    "Python, Go, PostgreSQL, Kafka, Kubernetes, Terraform",
]


def icon_font_pdf() -> bytes:
    """Contact line prefixed with icon-font glyphs mapped into the Private Use Area."""
    to_unicode = (
        b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
        b"/CMapName /IconMap def 1 begincodespacerange <00> <FF> endcodespacerange\n"
        b"3 beginbfchar\n<01> <F0E0>\n<02> <F095>\n<03> <F08C>\nendbfchar\n"
        b"endcmap CMapName currentdict /CMap defineresource pop end end"
    )
    lines: list[tuple[str, int, int, str]] = [("F2", 72, 740, _pdf_literal(NAME))]
    y = 720
    for glyph, value in (("<01>", EMAIL), ("<02>", PHONE), ("<03>", LINKEDIN)):
        lines.append(("F1", 72, y, glyph))
        lines.append(("F2", 90, y, _pdf_literal(value)))
        y -= 16
    for text in _BODY_LINES:
        y -= 16
        lines.append(("F2", 72, y, _pdf_literal(text)))
    content = _text_ops(lines)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R /F2 7 0 R >> >> /Contents 4 0 R >>",
        _stream(content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /IconGlyphs "
        b"/FirstChar 1 /LastChar 3 /Widths [700 700 700] /ToUnicode 6 0 R >>",
        _stream(to_unicode),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    return _assemble_pdf(objects)


WATERMARK_TEXT = "CONFIDENTIAL DRAFT"


def rotated_watermark_pdf() -> bytes:
    """Single-column resume with a large, semi-transparent 45-degree watermark
    drawn inside a form XObject (as Chromium draws text with CSS opacity).

    The watermark crosses the body rows, so if its diagonal bounding box took
    part in row reconstruction it would merge into a body row.
    """
    lines: list[tuple[str, int, int, str]] = [
        ("F2", 72, 740, _pdf_literal(NAME)),
        ("F2", 72, 724, _pdf_literal(f"{EMAIL} | {PHONE} | {LINKEDIN}")),
    ]
    y = 710
    for text in _BODY_LINES:
        y -= 28
        lines.append(("F2", 72, y, _pdf_literal(text)))
    content = _text_ops(lines) + b"\nq /Wm Do Q"
    watermark = (
        f"q /G1 gs BT /F2 48 Tf 0.7071 0.7071 -0.7071 0.7071 150 400 Tm "
        f"{_pdf_literal(WATERMARK_TEXT)} Tj ET Q"
    ).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F2 5 0 R >> /XObject << /Wm 6 0 R >> >> /Contents 4 0 R >>",
        _stream(content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        _stream(
            watermark,
            b" /Type /XObject /Subtype /Form /BBox [0 0 612 792]"
            b" /Resources << /Font << /F2 5 0 R >> /ExtGState << /G1 7 0 R >> >>",
        ),
        b"<< /Type /ExtGState /ca 0.15 /CA 0.15 >>",
    ]
    return _assemble_pdf(objects)


def cid_glyphs_pdf() -> bytes:
    """Body text in a CID font with no ToUnicode map, so glyphs stay unmapped."""
    lines: list[tuple[str, int, int, str]] = [
        ("F2", 72, 740, _pdf_literal(NAME)),
        ("F2", 72, 724, _pdf_literal(f"{EMAIL} | {PHONE}")),
    ]
    y = 700
    for text in _BODY_LINES:
        codes = "".join(f"{ord(char) - 29:04X}" for char in text)
        lines.append(("F1", 72, y, f"<{codes}>"))
        y -= 16
    content = _text_ops(lines)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R /F2 7 0 R >> >> /Contents 4 0 R >>",
        _stream(content),
        b"<< /Type /Font /Subtype /Type0 /BaseFont /SyntheticSubset "
        b"/Encoding /Identity-H /DescendantFonts [6 0 R] >>",
        b"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /SyntheticSubset "
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
        b"/DW 500 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    return _assemble_pdf(objects)


def decompression_bomb_pdf() -> bytes:
    """A tiny PDF whose single content stream inflates to 20 MB."""
    payload = zlib.compress(b" " * (20 * 1024 * 1024), 9)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        _stream(payload, b" /Filter /FlateDecode"),
    ]
    return _assemble_pdf(objects)


def legacy_doc_stub() -> bytes:
    """A minimal compound-file container that passes the .doc header validation."""
    header = bytearray(512)
    header[0:8] = bytes.fromhex("D0CF11E0A1B11AE1")
    header[24:26] = (0x3E).to_bytes(2, "little")
    header[26:28] = (3).to_bytes(2, "little")
    header[28:30] = b"\xfe\xff"
    header[30:32] = (9).to_bytes(2, "little")
    header[32:34] = (6).to_bytes(2, "little")
    header[44:48] = (1).to_bytes(4, "little")
    header[48:52] = (1).to_bytes(4, "little")
    header[56:60] = (4096).to_bytes(4, "little")
    header[60:64] = (0xFFFFFFFE).to_bytes(4, "little")
    header[68:72] = (0xFFFFFFFE).to_bytes(4, "little")
    for offset in range(76, 512, 4):
        header[offset : offset + 4] = (0xFFFFFFFF).to_bytes(4, "little")
    header[76:80] = (0).to_bytes(4, "little")
    fat = bytearray(b"\xff" * 512)
    fat[0:4] = (0xFFFFFFFD).to_bytes(4, "little")
    fat[4:8] = (0xFFFFFFFE).to_bytes(4, "little")
    directory = bytes(512)
    return bytes(header) + bytes(fat) + directory


def _docx_bytes(build: Callable[[Document], None]) -> bytes:
    document = Document()
    build(document)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _add_body(document: Document) -> None:
    document.add_heading("Summary", level=2)
    document.add_paragraph(SUMMARY)
    document.add_heading("Experience", level=2)
    for job in EXPERIENCE:
        document.add_paragraph(f"{job['title']}, {job['company']} ({job['years']})")
        for bullet in job["bullets"]:  # type: ignore[union-attr]
            document.add_paragraph(str(bullet), style="List Bullet")
    document.add_heading("Education", level=2)
    document.add_paragraph(
        f"{EDUCATION['degree']}, {EDUCATION['institution']} ({EDUCATION['years']})"
    )
    document.add_heading("Skills", level=2)
    document.add_paragraph(SKILLS)


def clean_docx(document: Document) -> None:
    """Single-column DOCX with contact details in the body."""
    document.add_heading(NAME, level=1)
    document.add_paragraph(f"{EMAIL} | {PHONE} | {LINKEDIN} | {LOCATION}")
    _add_body(document)


def contact_in_header_docx(document: Document) -> None:
    """Contact details exist only in the page header; the body has year ranges.

    Year ranges share a phone number's shape, so the body deliberately holds
    them to prove a header-only phone is still reported.
    """
    header = document.sections[0].header
    header.paragraphs[0].text = f"{NAME} | {EMAIL} | {PHONE} | {LINKEDIN}"
    document.add_heading("Professional Profile", level=1)
    _add_body(document)
    document.add_heading("Volunteering", level=2)
    document.add_paragraph("Mentor, Harbor Code Club 2015 - 2019 2019 - 2021")


TEXT_BOX_TEXT = "Certified Kubernetes Administrator, 2022"
_TEXT_BOX_RUN = """<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps">
<w:drawing><wp:inline><wp:extent cx="2743200" cy="457200"/><wp:docPr id="1" name="Box"/>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>
</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>
<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p>
</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"""


def text_box_docx(document: Document) -> None:
    """Clean body plus one Word text box (DrawingML choice with a VML fallback)."""
    clean_docx(document)
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(_TEXT_BOX_RUN.format(text=TEXT_BOX_TEXT)))


def table_layout_docx(document: Document) -> None:
    """Experience laid out in a table grid, a common template pattern."""
    document.add_heading(NAME, level=1)
    document.add_paragraph(f"{EMAIL} | {PHONE} | {LINKEDIN}")
    document.add_heading("Experience", level=2)
    table = document.add_table(rows=len(EXPERIENCE), cols=3)
    for row, job in zip(table.rows, EXPERIENCE, strict=True):
        row.cells[0].text = str(job["years"])
        row.cells[1].text = f"{job['title']}\n{job['company']}"
        row.cells[2].text = "\n".join(str(item) for item in job["bullets"])  # type: ignore[union-attr]
    document.add_heading("Education", level=2)
    document.add_paragraph(f"{EDUCATION['degree']}, {EDUCATION['institution']}")
    document.add_heading("Skills", level=2)
    document.add_paragraph(SKILLS)


def two_column_section_docx(document: Document) -> None:
    """Word section formatted with two text columns."""
    section = document.sections[0]
    section.start_type = WD_SECTION.CONTINUOUS
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "720")
    clean_docx(document)


def _render_browser_fixtures(
    html_fixtures: dict[str, Callable[[], str]], selected: set[str]
) -> dict[str, bytes]:
    """Render the Chromium-backed fixtures (HTML pages and the image-only PDF)."""
    outputs: dict[str, bytes] = {}
    with sync_playwright() as playwright:
        # CHROMIUM_EXECUTABLE lets a machine reuse an already-installed build.
        executable = os.environ.get("CHROMIUM_EXECUTABLE") or None
        browser = playwright.chromium.launch(executable_path=executable)
        page = browser.new_page()
        for name, build in html_fixtures.items():
            outputs[name] = _render_pdf(page, build())
        if not selected or "image_only.pdf" in selected:
            outputs["image_only.pdf"] = _image_only_pdf(page)
        browser.close()
    return outputs


def main(selected: set[str]) -> None:
    """Write every fixture, or only the named ones when names are given."""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, bytes] = {
        "icon_font.pdf": icon_font_pdf(),
        "cid_glyphs.pdf": cid_glyphs_pdf(),
        "rotated_watermark.pdf": rotated_watermark_pdf(),
        "decompression_bomb.pdf": decompression_bomb_pdf(),
        "legacy.doc": legacy_doc_stub(),
        "clean.docx": _docx_bytes(clean_docx),
        "contact_in_header.docx": _docx_bytes(contact_in_header_docx),
        "table_layout.docx": _docx_bytes(table_layout_docx),
        "two_column_section.docx": _docx_bytes(two_column_section_docx),
        "text_box.docx": _docx_bytes(text_box_docx),
        "swiss_single_like.source.json": (
            json.dumps(swiss_single_source(), indent=2, sort_keys=True) + "\n"
        ).encode("utf-8"),
    }
    html_fixtures: dict[str, Callable[[], str]] = {
        "clean_single_column.pdf": clean_single_column_html,
        "swiss_single_like.pdf": swiss_single_like_html,
        "two_column.pdf": two_column_html,
        "balanced_two_column.pdf": balanced_two_column_html,
        "spanish.pdf": spanish_html,
        "twelve_pages.pdf": long_document_html,
    }
    html_fixtures = {
        name: build
        for name, build in html_fixtures.items()
        if not selected or name in selected
    }
    if html_fixtures or not selected or "image_only.pdf" in selected:
        outputs.update(_render_browser_fixtures(html_fixtures, selected))
    for name, data in sorted(outputs.items()):
        if selected and name not in selected:
            continue
        (FIXTURE_DIR / name).write_bytes(data)
        print(f"wrote {name} ({len(data)} bytes)")


if __name__ == "__main__":
    main(set(sys.argv[1:]))
