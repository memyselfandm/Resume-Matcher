"""Render the own-output ATS parse-check fixtures from the real resume templates.

The PDFs are produced by the real download route, ``GET /api/v1/resumes/{id}/pdf``
(headless Chromium printing the Next.js ``/print/resumes/[id]`` page), so the
default test suite checks the column detector and the round-trip against what
users actually download, without needing Chromium or the frontend.

Requirements (opt-in; re-run only when a template changes):

* the Next.js frontend running at ``FRONTEND_BASE_URL`` (default
  ``http://localhost:3000``), whose server-side data origin is the default
  ``http://127.0.0.1:8000``;
* Playwright's Chromium (``uv run playwright install chromium``);
* port 8000 free: this script serves the backend in-process on
  ``127.0.0.1:8000`` against a temporary ``DATA_DIR``, seeds one synthetic
  resume, and renders it.

Run from ``apps/backend``::

    uv run python ../../scripts/generate_ats_render_fixtures.py

Writes ``apps/backend/tests/fixtures/ats_parse/renders/``: ``renders.tar.xz``
with one PDF per template (default settings) plus ``swiss-single-es.pdf``
(``lang=es``), stored byte for byte; ``source.json`` (the ``processed_resume``
payload the print page rendered); and ``manifest.json`` (the settings of every
PDF, the platform, and the Chromium version: templates use system font stacks,
so glyph mapping can differ between macOS and Linux renders). Ideally
regenerate in the production image so the fixtures match what users get. It then runs
``POST /api/v1/resumes/{id}/parse-check`` with ``all_templates`` and prints each
template's verdicts.

The persona is fictional (reserved example domains, 555 phone number). It
covers the round-trip edge cases: a repeated employer, a bullet naming the
next entry's employer, HTML bullets, CSS-uppercased headings, a hidden custom
section, and a visible custom section.
"""

from __future__ import annotations

import argparse
import asyncio
import copy
import io
import json
import lzma
import os
import sys
import tarfile
import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = REPO_ROOT / "apps" / "backend"
OUTPUT_DIR = BACKEND_DIR / "tests" / "fixtures" / "ats_parse" / "renders"
BACKEND_HOST = "127.0.0.1"
BACKEND_PORT = 8000
ARCHIVE_NAME = "renders.tar.xz"
# Per request against a deployed instance (a render is bounded server-side).
RENDER_TIMEOUT_SECONDS = 300.0
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

PERSONA: dict[str, Any] = {
    "personalInfo": {
        "name": "Jordan Rivera",
        "title": "Senior Backend Engineer",
        "email": "jordan.rivera@example.com",
        "phone": "(555) 010-4477",
        "location": "Oakland, CA",
        "website": "https://jordanrivera.example.com",
        "linkedin": "linkedin.com/in/jordan-rivera-example",
        "github": "github.com/jordan-rivera-example",
    },
    "summary": (
        "Backend engineer with 9 years of experience building reliable data platforms, "
        "developer tooling, and logistics systems. Focused on measurable reliability and "
        "delivery improvements."
    ),
    "workExperience": [
        {
            "id": 1,
            "title": "Senior Software Engineer",
            "company": "Northwind Analytics",
            "location": "San Francisco, CA",
            "years": "Jan 2021 - Present",
            "description": [
                "<strong>Led</strong> migration of 40 services to a shared deployment platform, "
                "reducing release time by 35%.",
                "Designed an event pipeline processing 2,000,000 records per day with "
                "<em>99.9%</em> availability.",
                "Shipped routing features for Northwind Analytics Maps used by 12 regional teams.",
            ],
        },
        {
            "id": 2,
            "title": "Software Engineer",
            "company": "Northwind Analytics",
            "location": "Oakland, CA",
            "years": "Jun 2018 - Dec 2020",
            "description": [
                "Built route optimization services used by 300 dispatchers across 12 hubs.",
                "Improved query latency by 60% by redesigning the reporting schema.",
                "Implemented contract tests that reduced production incidents by 25% &amp; "
                "cut on-call pages in half.",
            ],
        },
        {
            "id": 3,
            "title": "Junior Developer",
            "company": "Harbor Point Media",
            "location": "Vallejo, CA",
            "years": "Aug 2015 - May 2018",
            "description": [
                "Developed content publishing tools for a newsroom of 80 editors.",
                "Streamlined image processing jobs, cutting storage costs by 18%.",
            ],
        },
    ],
    "education": [
        {
            "id": 1,
            "institution": "Bay State University",
            "degree": "B.S. Computer Science",
            "years": "Aug 2011 - May 2015",
            "description": "Graduated with honors.",
        }
    ],
    "personalProjects": [
        {
            "id": 1,
            "name": "Tidewater Scheduler",
            "role": "Creator and Maintainer",
            "years": "Mar 2022 - Present",
            "description": [
                "Open-source job scheduler with 1,200 GitHub stars and 40 contributors.",
            ],
        }
    ],
    "additional": {
        "technicalSkills": [
            "Python",
            "Go",
            "PostgreSQL",
            "Kafka",
            "Kubernetes",
            "Terraform",
            "FastAPI",
        ],
        "languages": ["English (Native)", "Spanish (Professional)"],
        "certificationsTraining": ["Certified Kubernetes Administrator"],
        "awards": ["Northwind Engineering Excellence Award 2023"],
    },
    "customSections": {
        "publications": {
            "sectionType": "itemList",
            "items": [
                {
                    "id": 1,
                    "title": "Designing Reliable Event Pipelines",
                    "subtitle": "Bay Area Systems Journal",
                    "years": "2023",
                    "description": ["Case study of exactly-once delivery at scale."],
                }
            ],
        },
        "volunteering": {
            "sectionType": "stringList",
            "strings": ["Code mentor at Oakland Youth Coding Club"],
        },
    },
}

CUSTOM_SECTION_META = [
    {
        "id": "publications",
        "key": "publications",
        "displayName": "Publications",
        "sectionType": "itemList",
        "isDefault": False,
        "isVisible": True,
        "order": 6,
    },
    {
        "id": "volunteering",
        "key": "volunteering",
        "displayName": "Volunteering",
        "sectionType": "stringList",
        "isDefault": False,
        "isVisible": False,
        "order": 7,
    },
]

TEMPLATES = (
    "swiss-single",
    "swiss-two-column",
    "modern",
    "modern-two-column",
    "latex",
    "clean",
    "vivid",
)
# Extra renders of swiss-single: (file stem, settings overrides).
VARIANTS: tuple[tuple[str, dict[str, Any]], ...] = (("swiss-single-es", {"lang": "es"}),)


def _persona() -> dict[str, Any]:
    from app.schemas.models import DEFAULT_SECTION_META

    data = copy.deepcopy(PERSONA)
    meta = [
        {**entry, "sectionType": str(getattr(entry["sectionType"], "value", entry["sectionType"]))}
        for entry in copy.deepcopy(DEFAULT_SECTION_META)
    ]
    data["sectionMeta"] = meta + copy.deepcopy(CUSTOM_SECTION_META)
    return data


async def _serve_backend() -> tuple[Any, asyncio.Task[None]]:
    import uvicorn

    from app.main import app

    server = uvicorn.Server(
        uvicorn.Config(app, host=BACKEND_HOST, port=BACKEND_PORT, lifespan="off", log_level="warning")
    )
    task = asyncio.create_task(server.serve())
    for _ in range(100):
        if server.started:
            return server, task
        if task.done():
            task.result()
        await asyncio.sleep(0.05)
    raise RuntimeError(f"Backend did not start on {BACKEND_HOST}:{BACKEND_PORT}")


async def _render_fixtures(client: Any, resume_id: str, describe: Callable[[], dict[str, Any]]) -> None:
    """Render every fixture through the download route and write the fixture files.

    ``client`` is an ``httpx.AsyncClient`` bound to the backend API (in-process
    or a deployed instance); ``describe`` returns where the renders ran and
    is called after the renders, once the renderer is known.
    """
    from app.services.ats_parse.own_output import TemplateSettings

    payload = (await client.get("/api/v1/resumes", params={"resume_id": resume_id})).json()
    source = payload["data"]["processed_resume"]
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "source.json").write_text(
        json.dumps(source, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    )
    jobs = [(template, {"template": template}) for template in TEMPLATES] + [
        (stem, {"template": "swiss-single", **overrides}) for stem, overrides in VARIANTS
    ]
    manifest: dict[str, dict[str, Any]] = {}
    renders: dict[str, bytes] = {}
    for stem, overrides in jobs:
        settings = TemplateSettings.model_validate(overrides)
        response = await client.get(
            f"/api/v1/resumes/{resume_id}/pdf",
            params=settings.pdf_query(settings.template),
        )
        response.raise_for_status()
        renders[f"{stem}.pdf"] = response.content
        manifest[stem] = settings.model_dump(mode="json")
        print(f"rendered {stem}.pdf ({len(response.content)} bytes)")
    _write_archive(renders)
    # Rendered text depends on the renderer host's fonts (the templates use
    # ui-sans-serif/system-ui stacks), so record where it ran.
    (OUTPUT_DIR / "manifest.json").write_text(
        json.dumps({**describe(), "files": manifest}, indent=2, sort_keys=True) + "\n"
    )
    response = await client.post(
        f"/api/v1/resumes/{resume_id}/parse-check",
        json={"all_templates": True, "content_language": "en"},
    )
    response.raise_for_status()
    _print_summary(response.json())


async def _generate_local() -> None:
    import httpx

    from app.database import db
    from app.main import app
    from app.pdf import close_pdf_renderer

    persona = _persona()
    resume = await db.create_resume(
        content=json.dumps(persona),
        content_type="json",
        processed_data=persona,
        processing_status="ready",
        is_master=True,
    )
    server, task = await _serve_backend()
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://fixtures.local", timeout=None
        ) as client:
            await _render_fixtures(
                client,
                resume["resume_id"],
                lambda: {"platform": sys.platform, "chromium": _chromium_version()},
            )
    finally:
        server.should_exit = True
        await task
        await close_pdf_renderer()
        await db.close()


async def _generate_remote(base_url: str, platform: str, chromium: str, image_id: str) -> None:
    """Seed the persona into a deployed instance through its REST API and render there.

    The upload only creates the resume (its LLM parse may fail without a
    configured provider); ``PATCH /resumes/{id}`` then stores the persona as
    its structured data, exactly as the builder saves it.
    """
    import httpx

    async with httpx.AsyncClient(base_url=base_url, timeout=RENDER_TIMEOUT_SECONDS) as client:
        upload = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("persona.docx", _persona_docx(), DOCX_MIME)},
        )
        upload.raise_for_status()
        resume_id = upload.json()["resume_id"]
        saved = await client.patch(f"/api/v1/resumes/{resume_id}", json=_persona())
        saved.raise_for_status()
        await _render_fixtures(
            client,
            resume_id,
            lambda: {"platform": platform, "chromium": chromium, "image_id": image_id},
        )


def _persona_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(PERSONA["personalInfo"]["name"])
    document.add_paragraph(PERSONA["summary"])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _chromium_version() -> str | None:
    from app import pdf

    browser = getattr(pdf, "_browser", None)
    return browser.version if browser is not None else None


def _write_archive(renders: dict[str, bytes]) -> None:
    """Store the PDFs byte for byte in one xz tarball.

    The seven templates embed largely the same font subsets, so one archive
    compresses them to a fraction of their size; members carry a fixed mtime
    so regenerating identical PDFs yields an identical archive.
    """
    buffer = io.BytesIO()
    with tarfile.open(fileobj=buffer, mode="w:xz", preset=9 | lzma.PRESET_EXTREME) as archive:
        for name in sorted(renders):
            info = tarfile.TarInfo(name)
            info.size = len(renders[name])
            info.mtime = 0
            archive.addfile(info, io.BytesIO(renders[name]))
    (OUTPUT_DIR / ARCHIVE_NAME).write_bytes(buffer.getvalue())
    print(f"wrote {ARCHIVE_NAME} ({len(buffer.getvalue())} bytes)")


def _print_summary(body: dict[str, Any]) -> None:
    print("template | status | multi_column | sidebar | expected | recall | order | overall | content")
    for result in body["results"]:
        report = result.get("report") or {}
        checks = {check["id"]: check for check in report.get("checks", [])}
        roundtrip = report.get("roundtrip") or {}
        print(
            " | ".join(
                str(value)
                for value in (
                    result["template"],
                    result["status"],
                    checks.get("multi_column", {}).get("status"),
                    checks.get("sidebar", {}).get("status"),
                    result["expected_by_template"],
                    roundtrip.get("content_recall"),
                    roundtrip.get("order_fidelity"),
                    report.get("overall_score"),
                    report.get("content_score"),
                )
            )
        )


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    parser.add_argument(
        "--base-url",
        help="render on a deployed instance (e.g. the production Docker image) instead of "
        "serving the backend in-process",
    )
    parser.add_argument("--platform", default="linux", help="with --base-url: the renderer's OS")
    parser.add_argument("--chromium", help="with --base-url: the Chromium version it runs")
    parser.add_argument("--image-id", help="with --base-url: the Docker image ID")
    args = parser.parse_args(argv)
    if args.base_url and not (args.chromium and args.image_id):
        parser.error("--base-url needs --chromium and --image-id for the manifest")
    return args


def main() -> None:
    args = _parse_args()
    os.chdir(BACKEND_DIR)
    sys.path.insert(0, str(BACKEND_DIR))
    with tempfile.TemporaryDirectory(prefix="rm-render-fixtures-") as data_dir:
        # DATA_DIR must be set before the app (and its database) is imported.
        os.environ["DATA_DIR"] = data_dir
        if args.base_url:
            asyncio.run(
                _generate_remote(args.base_url, args.platform, args.chromium, args.image_id)
            )
        else:
            asyncio.run(_generate_local())


if __name__ == "__main__":
    main()
